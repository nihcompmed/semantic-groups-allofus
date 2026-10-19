"""workbook.py -- what the supplementary files and the figure source data share: the Excel writer, the Holm adjustment,
the number formats, and the names of the groupings.

Every workbook opens with a "Read me" sheet (what each sheet holds and how to read its columns), then one sheet per
table. Each data sheet is also written as CSV to elife/supplementary/csv/<workbook>__<sheet>.csv.
"""
import os

import numpy as np

from common import ENCODERS, PRIMARY, SUPPLEMENTARY, display_profile

OTHER = ["tfidf4", "tfidf25", "pca", "factor", "mca90", "mca"]   # the 6 groupings that do not use a sentence encoder
GLAB = {PRIMARY: "gte-large-en-v1.5 (prespecified)", "bge-m3": "bge-m3", "bge-large-en-v1.5": "bge-large-en-v1.5",
        "all-mpnet-base-v2": "all-mpnet-base-v2", "e5-large-v2": "e5-large-v2", "qwen3-embedding-0.6b": "Qwen3-Embedding-0.6B",
        "tfidf4": "Word counts (TF-IDF), Horn's k", "tfidf25": "Word counts (TF-IDF), 25 groups",
        "pca": "Principal components of the answers", "factor": "Factor model of the answers",
        "mca90": "Multiple correspondence analysis, Horn's k", "mca": "Multiple correspondence analysis, 25 groups",
        "cohort": "Whole population (no decile)", "population": "Whole population (no decile)"}
GTYPE = {**{e: "sentence encoder" for e in ENCODERS}, "tfidf4": "item text", "tfidf25": "item text", "pca": "answers",
         "factor": "answers", "mca90": "answers", "mca": "answers"}


def holm(p):
    """Holm-adjusted p values, in the order given."""
    p = np.asarray(p, float)
    m, out, run = len(p), np.empty(len(p)), 0.0
    for i, ix in enumerate(np.argsort(p)):
        run = max(run, min(1.0, (m - i) * p[ix]))
        out[ix] = run
    return out


def pv(p):
    """P for a sheet: 3 significant digits; a P that underflows to 0 in double precision is written < 1e-300."""
    return "< 1e-300" if p == 0 else float(f"{p:.3g}")


def npct(n, pct):
    """No. (%) with a thousands comma; a count already written as text (e.g. <=20) is kept."""
    try:
        return f"{int(float(n)):,} ({pct:.1f})"
    except (TypeError, ValueError):
        return f"{n} ({pct:.1f})"


def names(profile):
    """The constructs of a unit's core items with their counts, as the figures print them ("Loneliness 6, ...")."""
    return ", ".join(x.replace(":", " ") for x in display_profile(profile).split("; "))


def write(name, readme, sheets, widths=(30, 130), csv=True):
    """elife/supplementary/<name>.xlsx: the Read me sheet, then each (sheet name, data frame); a missing number is an
    empty cell."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    os.makedirs(os.path.join(SUPPLEMENTARY, "csv"), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for k, v in readme:
        ws.append([k, v])
    for row in ws.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = widths[0]
    ws.column_dimensions["B"].width = widths[1]
    for sname, df in sheets:
        s = wb.create_sheet(sname[:31])
        s.append(list(df.columns))
        for c in s[1]:
            c.font = Font(bold=True)
        for row in df.itertuples(index=False):
            s.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
        s.freeze_panes = "A2"
        if csv:
            df.to_csv(os.path.join(SUPPLEMENTARY, "csv", f"{name}__{sname.replace(' ', '_')}.csv"), index=False)
    wb.save(os.path.join(SUPPLEMENTARY, f"{name}.xlsx"))
    print(f"-> {name}.xlsx: " + ", ".join(f"{s} {len(d)}" for s, d in sheets))
